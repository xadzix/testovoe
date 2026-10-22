from collections.abc import Iterator
from datetime import date, datetime, time, timedelta
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from django.conf import settings
from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from apps.data_imports.contracts import (
    CellPresence,
    ExcelCellValue,
    FormulaResultStatus,
    XlsxSourceCell,
    XlsxSourceRow,
)


class InvalidXlsxFileError(ValueError):
    """Исходный файл нельзя открыть как корректную книгу XLSX."""


class WorksheetNotFoundError(ValueError):
    """Запрошенный лист отсутствует в книге."""


class UnsupportedCellValueError(TypeError):
    """openpyxl вернул значение вне поддерживаемого набора скалярных типов."""


class XlsxSourceReader:
    """Последовательно читать строки, не загружая лист целиком в память."""

    def __init__(self, source_path: str | Path) -> None:
        self._source_path = Path(source_path)

    def iter_rows(self, sheet_name: str | None = None) -> Iterator[XlsxSourceRow]:
        self._validate_archive()
        formula_workbook = self._open_workbook(data_only=False)
        try:
            cached_workbook = self._open_workbook(data_only=True)
        except InvalidXlsxFileError:
            formula_workbook.close()
            raise

        try:
            selected_sheets = self._select_sheets(
                formula_workbook.worksheets,
                cached_workbook.worksheets,
                sheet_name,
            )
            for formula_sheet, cached_sheet in selected_sheets:
                yield from self._iter_sheet_rows(
                    formula_sheet,
                    cached_sheet,
                    formula_workbook.epoch,
                )
        finally:
            formula_workbook.close()
            cached_workbook.close()

    def _validate_archive(self) -> None:
        try:
            with ZipFile(self._source_path) as archive:
                entries = archive.infolist()
        except (BadZipFile, OSError, ValueError) as error:
            raise InvalidXlsxFileError("The uploaded file is not a valid XLSX workbook") from error

        if len(entries) > int(settings.XLSX_MAX_ARCHIVE_ENTRIES):
            raise InvalidXlsxFileError("The XLSX archive exceeds configured safety limits")

        total_size = 0
        max_part_size = int(settings.XLSX_MAX_PART_SIZE_BYTES)
        max_total_size = int(settings.XLSX_MAX_UNCOMPRESSED_SIZE_BYTES)
        for entry in entries:
            if entry.file_size > max_part_size:
                raise InvalidXlsxFileError("The XLSX archive exceeds configured safety limits")
            total_size += entry.file_size
            if total_size > max_total_size:
                raise InvalidXlsxFileError("The XLSX archive exceeds configured safety limits")

    def _open_workbook(self, *, data_only: bool) -> Workbook:
        try:
            return load_workbook(
                self._source_path,
                read_only=True,
                data_only=data_only,
                keep_links=False,
            )
        except (
            BadZipFile,
            DefusedXmlException,
            InvalidFileException,
            KeyError,
            OSError,
            ParseError,
            ValueError,
        ) as error:
            raise InvalidXlsxFileError("The uploaded file is not a valid XLSX workbook") from error

    @staticmethod
    def _select_sheets(
        formula_sheets: list[ReadOnlyWorksheet],
        cached_sheets: list[ReadOnlyWorksheet],
        sheet_name: str | None,
    ) -> tuple[tuple[ReadOnlyWorksheet, ReadOnlyWorksheet], ...]:
        cached_by_name = {sheet.title: sheet for sheet in cached_sheets}
        formula_by_name = {sheet.title: sheet for sheet in formula_sheets}

        if sheet_name is not None:
            if sheet_name not in formula_by_name:
                raise WorksheetNotFoundError(f"Worksheet {sheet_name!r} does not exist")
            return ((formula_by_name[sheet_name], cached_by_name[sheet_name]),)

        return tuple((sheet, cached_by_name[sheet.title]) for sheet in formula_sheets)

    def _iter_sheet_rows(
        self,
        formula_sheet: ReadOnlyWorksheet,
        cached_sheet: ReadOnlyWorksheet,
        workbook_epoch: datetime,
    ) -> Iterator[XlsxSourceRow]:
        first_row_number = formula_sheet.min_row or 1
        paired_rows = zip(
            formula_sheet.iter_rows(),
            cached_sheet.iter_rows(),
            strict=True,
        )
        for row_number, (formula_row, cached_row) in enumerate(
            paired_rows,
            start=first_row_number,
        ):
            cells = tuple(
                self._build_cell(
                    formula_cell,
                    cached_cell,
                    formula_sheet.title,
                    row_number,
                    column_number,
                )
                for column_number, (formula_cell, cached_cell) in enumerate(
                    zip(formula_row, cached_row, strict=True),
                    start=1,
                )
            )
            yield XlsxSourceRow(formula_sheet.title, row_number, cells, workbook_epoch)

    @staticmethod
    def _build_cell(
        formula_cell: object,
        cached_cell: object,
        sheet_name: str,
        row_number: int,
        column_number: int,
    ) -> XlsxSourceCell:
        coordinate = f"{get_column_letter(column_number)}{row_number}"
        if isinstance(formula_cell, EmptyCell):
            return XlsxSourceCell(
                sheet_name=sheet_name,
                coordinate=coordinate,
                row_number=row_number,
                column_number=column_number,
                raw_value=None,
                data_type="n",
                number_format="General",
                is_date=False,
                presence=CellPresence.MISSING,
                formula_expression=None,
                formula_cached_value=None,
                formula_result_status=FormulaResultStatus.NOT_FORMULA,
            )

        raw_value = _as_excel_value(getattr(formula_cell, "value", None))
        data_type = str(getattr(formula_cell, "data_type", "n"))
        formula_expression = raw_value if data_type == "f" and isinstance(raw_value, str) else None
        cached_value = _as_excel_value(getattr(cached_cell, "value", None))

        if formula_expression is None:
            formula_status = FormulaResultStatus.NOT_FORMULA
            cached_value = None
        elif cached_value is None:
            formula_status = FormulaResultStatus.MISSING
        else:
            formula_status = FormulaResultStatus.AVAILABLE

        presence = CellPresence.VALUE if raw_value is not None else CellPresence.EXPLICIT_EMPTY
        return XlsxSourceCell(
            sheet_name=sheet_name,
            coordinate=coordinate,
            row_number=row_number,
            column_number=column_number,
            raw_value=raw_value,
            data_type=data_type,
            number_format=str(getattr(formula_cell, "number_format", "General")),
            is_date=bool(getattr(formula_cell, "is_date", False)),
            presence=presence,
            formula_expression=formula_expression,
            formula_cached_value=cached_value,
            formula_result_status=formula_status,
        )


def _as_excel_value(value: object) -> ExcelCellValue:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value

    if isinstance(value, (date, datetime, time, timedelta)):
        return value
    raise UnsupportedCellValueError(f"Unsupported XLSX cell value type: {type(value).__name__}")
