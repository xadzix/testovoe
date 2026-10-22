from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from django.conf import LazySettings
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.worksheet import Worksheet

from apps.data_imports.contracts import CellPresence, FormulaResultStatus
from apps.data_imports.readers.xlsx import InvalidXlsxFileError, XlsxSourceReader

STANDARD_HEADERS = (
    "date of request",
    "mark",
    "catalog number",
    "product name",
    "brand",
    "quantity",
    "price",
    "delivery time",
    "original / analogue",
)


def _active_sheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Requests"
    return sheet


def _save_workbook(workbook: Workbook, path: Path) -> Path:
    workbook.save(path)
    workbook.close()
    return path


def test_reads_rows_sequentially_with_source_metadata(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.append(STANDARD_HEADERS)
    sheet.append((date(2026, 10, 20), "ACME", "001-AB", "Bearing", "SKF", 2, 10.5, 5, "original"))
    sheet.append((date(2026, 10, 20), "ACME", "001-AB", "Bearing", "SKF", 2, 10.5, 5, "original"))
    source_path = _save_workbook(workbook, tmp_path / "requests.xlsx")

    rows = XlsxSourceReader(source_path).iter_rows("Requests")
    header_row = next(rows)
    first_item = next(rows)
    repeated_item = next(rows)

    assert header_row.sheet_name == "Requests"
    assert header_row.row_number == 1
    assert first_item.row_number == 2
    assert first_item.cells[0].coordinate == "A2"
    assert first_item.cells[0].raw_value == datetime(2026, 10, 20)
    assert first_item.cells[0].is_date is True
    assert first_item.cells[2].raw_value == "001-AB"
    assert first_item.cells[5].raw_value == 2
    assert first_item.cells[6].raw_value == 10.5
    assert repeated_item.cells[1].raw_value == "ACME"


def test_distinguishes_missing_and_explicit_empty_cells(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet["B1"] = None
    sheet["B1"].number_format = "0.00"
    sheet["C1"] = "keeps row width"
    sheet["A3"] = "after empty row"
    source_path = _save_workbook(workbook, tmp_path / "empty-cells.xlsx")

    rows = XlsxSourceReader(source_path).iter_rows()
    first_row = next(rows)
    empty_row = next(rows)

    assert first_row.cells[0].presence is CellPresence.MISSING
    assert first_row.cells[1].presence is CellPresence.EXPLICIT_EMPTY
    assert first_row.cells[1].number_format == "0.00"
    assert empty_row.is_empty is True


def test_preserves_number_formats_and_native_numeric_types(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet["A1"] = 7
    sheet["A1"].number_format = "000000"
    sheet["B1"] = 12.375
    sheet["C1"] = date(2026, 1, 15)
    source_path = _save_workbook(workbook, tmp_path / "typed-values.xlsx")

    row = next(XlsxSourceReader(source_path).iter_rows())

    assert row.cells[0].raw_value == 7
    assert row.cells[0].number_format == "000000"
    assert row.cells[1].raw_value == 12.375
    assert row.cells[2].raw_value == datetime(2026, 1, 15)
    assert row.cells[2].is_date is True


def test_preserves_workbook_epoch_for_numeric_date_conversion(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    sheet = _active_sheet(workbook)
    sheet["A1"] = 42_000
    source_path = _save_workbook(workbook, tmp_path / "mac-epoch.xlsx")

    row = next(XlsxSourceReader(source_path).iter_rows())

    assert row.workbook_epoch == CALENDAR_MAC_1904
    assert row.cells[0].raw_value == 42_000


def test_reports_formula_without_cached_value(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet["A1"] = 10
    sheet["B1"] = "=A1*2"
    source_path = _save_workbook(workbook, tmp_path / "formula.xlsx")

    formula_cell = next(XlsxSourceReader(source_path).iter_rows()).cells[1]

    assert formula_cell.raw_value == "=A1*2"
    assert formula_cell.formula_expression == "=A1*2"
    assert formula_cell.formula_cached_value is None
    assert formula_cell.formula_result_status is FormulaResultStatus.MISSING
    assert formula_cell.value_for_conversion is None


def test_merged_continuation_is_indistinguishable_from_missing_cell(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet["A1"] = "customer"
    sheet.merge_cells("A1:A2")
    sheet["B2"] = "item"
    source_path = _save_workbook(workbook, tmp_path / "merged.xlsx")

    rows = XlsxSourceReader(source_path).iter_rows()
    anchor_row = next(rows)
    continuation_row = next(rows)

    assert anchor_row.cells[0].raw_value == "customer"
    assert anchor_row.cells[0].presence is CellPresence.VALUE
    assert continuation_row.cells[0].raw_value is None
    assert continuation_row.cells[0].presence is CellPresence.MISSING


def test_rejects_corrupted_xlsx(tmp_path: Path) -> None:
    source_path = tmp_path / "corrupted.xlsx"
    source_path.write_bytes(b"not an XLSX archive")

    with pytest.raises(InvalidXlsxFileError):
        next(XlsxSourceReader(source_path).iter_rows())


@pytest.mark.parametrize(
    ("setting_name", "limit"),
    [
        ("XLSX_MAX_ARCHIVE_ENTRIES", 2),
        ("XLSX_MAX_UNCOMPRESSED_SIZE_BYTES", 29),
        ("XLSX_MAX_PART_SIZE_BYTES", 9),
    ],
)
def test_rejects_archive_exceeding_safety_limits(
    setting_name: str,
    limit: int,
    tmp_path: Path,
    settings: LazySettings,
) -> None:
    source_path = tmp_path / "oversized-archive.xlsx"
    with ZipFile(source_path, "w") as archive:
        for index in range(3):
            archive.writestr(f"part-{index}.xml", b"1234567890")
    setattr(settings, setting_name, limit)

    with pytest.raises(InvalidXlsxFileError, match="configured safety limits"):
        next(XlsxSourceReader(source_path).iter_rows())


def test_streams_large_worksheet_without_returning_a_collection(tmp_path: Path) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Requests")
    sheet.append(STANDARD_HEADERS)
    for row_number in range(1, 5_001):
        sheet.append(
            (None, None, f"PART-{row_number:06d}", "Part", "Brand", 1, 1.25, 3, "analogue")
        )
    source_path = _save_workbook(workbook, tmp_path / "large.xlsx")

    rows = XlsxSourceReader(source_path).iter_rows("Requests")

    assert iter(rows) is rows
    assert sum(1 for _row in rows) == 5_001
